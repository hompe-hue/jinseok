"""Refresh existing workbook audit sheets and regenerate guide from Markdown."""
from pathlib import Path
import json,pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Pt
root=Path(__file__).resolve().parents[1]
wb=load_workbook(root/'Bolt_Validation_Inputs_and_Audit_20261006.xlsx')
for name,file in [('Voltage_residuals','battery_voltage_by_SOC.csv'),('SOC_channel_fit','display_to_internal_SOC_observed_fit.csv')]:
 w=wb[name];w.delete_rows(1,w.max_row);d=pd.read_csv(root/'results'/file);w.append(list(d.columns))
 for r in d.itertuples(index=False,name=None):w.append(list(r))
audit=json.loads((root/'results/battery_data_quality.json').read_text());w=wb['Battery_data_quality'];w.delete_rows(1,w.max_row)
cols=['test','native_phase_rows','unique_meter_rows','invalid_voltage_or_current_rows','integrated_valid_seconds','phase_declared_seconds','meter_phase_span_s','coverage_fraction','integrated_duration_to_declared_phase_duration_ratio','integrated_valid_Ah','integrated_valid_meter_Wh','declared_phase_Ah','declared_phase_Wh','temperature_min_degC','temperature_max_degC'];w.append(cols)
for r in audit:w.append([r[c] for c in cols])
w=wb['Apparent_capacity'];w.delete_rows(1,w.max_row);cols=['test','channel','initial_SOC_percent','terminal_SOC_percent','valid_segment_net_Ah','apparent_Ah_per_unit_channel_SOC','replayed_terminal_SOC_percent_assuming_168Ah','note'];w.append(cols)
for r in audit:
 for c in r['apparent_endpoint_capacity_not_cell_capacity']:w.append([r['test']]+[c[k] for k in cols[1:]])
wb['READ_ME'].append(['Analysis voltage screen','200-450V data-plausibility screen, not a BMS limit. Physical low-voltage discharge tail retained.'])
for w in wb:
 w.freeze_panes='A2';w.auto_filter.ref=w.dimensions
 for cell in w[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234E70')
 for row in w:
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
wb.save(root/'Bolt_Validation_Inputs_and_Audit_20261006.xlsx')
doc=Document();doc.styles['Normal'].font.name='Malgun Gothic';doc.styles['Normal'].font.size=Pt(10)
lines=(root/'README_검증보완.md').read_text().splitlines();i=0
while i<len(lines):
 line=lines[i]
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   cells=[c.strip() for c in lines[i].strip('|').split('|')]
   if not all(set(c)<=set('-: ') for c in cells):rows.append(cells)
   i+=1
  table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
  for r in rows:
   for cell,value in zip(table.add_row().cells,r):cell.text=value
  continue
 if line.startswith('# '):doc.add_heading(line[2:],0)
 elif line.startswith('## '):doc.add_heading(line[3:],1)
 elif line.startswith('### '):doc.add_heading(line[4:],2)
 elif line and not line.startswith('```'):doc.add_paragraph(line.replace('**','').replace('`',''))
 i+=1
doc.save(root/'Bolt_Validation_Guide_20261006.docx')
print('Workbook and guide refreshed from final audit')
