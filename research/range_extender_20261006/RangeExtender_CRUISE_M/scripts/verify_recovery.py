"""Verify preserved archive members and all workbook CSV tables without rebuilding maps."""
from pathlib import Path
import csv
import hashlib
import json
import math
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT.parent
verified = []
for item in json.loads((ROOT / 'logs/deliverable_manifest.json').read_text()):
    path = ROOT / item['file']
    assert path.is_file(), item['file']
    data = path.read_bytes()
    assert len(data) == item['bytes'], item['file']
    assert hashlib.sha256(data).hexdigest() == item['sha256'], item['file']
    verified.append(item['file'])

wb = load_workbook(WORK / 'RangeExtender_CRUISE_M_Inputs_20261006.xlsx', read_only=True, data_only=True)
matched = []
for row in list(wb['INDEX'].values)[1:]:
    sheet, filename = row[0], row[1]
    if not filename or not sheet:
        continue
    path = ROOT / 'inputs' / filename
    with path.open(encoding='utf-8-sig', newline='') as stream:
        csv_rows = list(csv.reader(stream))
    xlsx_rows = list(wb[sheet].values)
    assert len(csv_rows) == len(xlsx_rows), filename
    for a, b in zip(csv_rows, xlsx_rows):
        assert len(a) == len(b), filename
        for x, y in zip(a, b):
            if y is None:
                assert x == '', (filename, x, y)
            elif isinstance(y, (int, float)):
                assert math.isclose(float(x), float(y), rel_tol=1e-12, abs_tol=1e-12), (filename, x, y)
            else:
                assert x == y, (filename, x, y)
    matched.append(filename)

assert len(matched) == 41
report = {
    'preserved_manifest_members': len(verified),
    'all_preserved_hashes_match': True,
    'csv_tables_equal_workbook': matched,
    'csv_table_count': len(matched),
    'workbook_sheets': len(wb.sheetnames),
    'maps_rebuilt': False,
    'CRUISE_M_GUI_import_verified': False,
    'CRUISE_M_executed': False,
}
(ROOT / 'logs/recovery_verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps({k: v for k, v in report.items() if k != 'csv_tables_equal_workbook'}))
