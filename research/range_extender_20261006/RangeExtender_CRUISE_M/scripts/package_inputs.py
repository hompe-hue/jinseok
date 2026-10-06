"""R65C-03: produce browsable spreadsheet, guide, and reproducible CSV kit."""
from pathlib import Path
import csv,json,hashlib,zipfile,re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml.ns import qn

R=Path(__file__).resolve().parents[1]; OUT=R.parents[1]/'output'; IN=R/'inputs'
OUT.mkdir(exist_ok=True)
def read_csv(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.reader(f))
def num(v):
    if v=='':return None
    try:return float(v)
    except ValueError:return v
def scalar_rows():
    return [
      ['Engine','Engine type','Reciprocating 4 stroke','Reciprocating 4 stroke','UI/map classification'],
      ['Engine','Transient model','Naturally aspirated','Naturally aspirated','NA; transient calibration missing'],
      ['Engine','Control mode','Torque controlled','Torque controlled','speed loop not supplied'],
      ['Engine','Fuel consumption map unit','Expressed as mass flow','Expressed as mass flow','Check actual Program g/s or kg/h'],
      ['Engine','Displacement [L]',1.998,1.5,'REF source; A assumed'],
      ['Engine','Number of cylinders',4,4,'REF source; A retained'],
      ['Engine','Engine inertia [kg*m2]',0.134,0.134,'USER SCREENSHOT DEFAULT; not target measurement'],
      ['Engine','Auxiliary inertia [kg*m2]',0,0,'USER SCREENSHOT DEFAULT; coupling unresolved'],
      ['Engine','Stall speed [rpm]',500,500,'USER SCREENSHOT DEFAULT; not target measurement'],
      ['Engine','Maximum speed [rpm]',6875,6875,'REF envelope endpoint; A retained, endpoint torque0'],
      ['Engine','Delta static friction [N*m]',0.1,0.1,'USER SCREENSHOT DEFAULT; not target measurement'],
      ['Engine','Specific carbon content [-]',0.826,0.826,'FTAG24350; A same reference fuel'],
      ['Engine','Fuel lower heat [kJ/kg]',41835,41835,'FTAG24350; A same reference fuel'],
      ['Engine','Fuel density [kg/m3]',750.2,750.2,'15C reference fuel; A retained'],
      ['Engine','Max shaft power [kW]',115.610609652,80,'Curve/controller cap; NOT continuous certified rating'],
      ['Generator','Type','IPM/PMSM','Virtual IPM/PMSM','Select available synchronous option; not current Asynchronous'],
      ['Generator','Global losses map','motor+inverter aggregate','motor+inverter aggregate','Match actual block boundary; do not duplicate inverter loss'],
      ['Generator','Characteristics expressed as','Torque','Torque','UI'],
      ['Generator','Losses depending on','Torque','Torque','UI'],
      ['Generator','Losses','Efficiency','Efficiency','Input percentage, not fraction'],
      ['Generator','Control mode','Mechanical load','Mechanical load','Q4 negative torque at positive rpm'],
      ['Generator','Quadrants','Q1 and Q4','Q1 and Q4','No reverse rotation charts'],
      ['Generator','Inertia MG1 [kg*m2]',0.00562277135625,0.05,'REF EPA parameter; A USER SCREENSHOT DEFAULT, no scaling law'],
      ['Generator','Inertia MG2 [kg*m2]',0.010773667,'not adopted','Alternative reference only'],
      ['Generator','Source maximum speed [rpm]',13500,13500,'Reference parameter, not map usability limit'],
      ['Generator','Use maximum speed [rpm]',13000,13000,'Guard: efficiency grid ends13000'],
      ['Generator','Minimum efficiency grid speed [rpm]',1000,1000,'Starting below1000rpm needs separate treatment'],
      ['Generator','Voltage coordinate [V]',650,650,'NOT330V direct-connected map'],
      ['Generator','Temperature coordinate [degC]',0,0,'INDEX ONLY, NOT0C test; no temperature dependence'],
      ['Generator','Shaft power cap [kW]',60,80,'Enforce controller cap, NOT certified continuous output'],
      ['Genset','Net bus cap [kW]','not65',65,'A: after converter and genset aux; controller parameter'],
      ['Genset','DCDC efficiency [-]','not adopted',0.97,'A assumption; separate from inverter already in efficiency map'],
      ['Genset','Genset auxiliary power [kW]','not adopted',1,'A assumption; not vehicle accessories'],
      ['Genset','Generator/engine speed ratio','not adopted',1,'A1:1 shaft coupling'],
      ['Genset','Coupling efficiency [-]','not adopted',1,'A lossless coupling assumption'],
      ['Genset','Cold start / delay / ramp','missing','missing','Do not assert measured dynamic response'],
    ]
with (IN/'Scalar_presets.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(['Component','Parameter','REF_value','A_virtual_value','Evidence_or_limit']);w.writerows(scalar_rows())

# Independent readback checks on exported tables, not source-only assertions.
checks=[]
for prefix in ['REF_MG1_EST','REF_MG2','A_Gen80_MG1EST']:
    rows=read_csv(IN/(prefix+'_Efficiency_Q1Q4.csv'))
    assert rows[0]==['Temperature (degC)','Voltage (V)','Speed (rpm)','Torque (N*m)','Efficiency (%)']
    a=np.asarray([[float(x) for x in row] for row in rows[1:]])
    assert a.shape==(494,5) and np.isfinite(a).all()
    assert np.all(a[:,1]==650) and np.all(a[:,2]>0) and np.all((a[:,4]>0)&(a[:,4]<=100))
    assert np.sum(a[:,3]>0)==247 and np.sum(a[:,3]<0)==247
    for q,sign in [('Q1',1),('Q4',-1)]:
        b=np.asarray([[float(x) for x in row] for row in read_csv(IN/(prefix+'_FullLoad_'+q+'.csv'))[1:]])
        assert np.all(b[:,0]==650) and np.all(np.diff(b[:,1])>0) and np.all(b[:,2]*sign>=0)
    checks.append(prefix+': 494 efficiency rows, signs and percentages verified')
for prefix in ['REF_Engine20NA','A_Engine15NA80']:
    a=np.asarray([[float(x) for x in row] for row in read_csv(IN/(prefix+'_Fuel_g_s.csv'))[1:]])
    b=np.asarray([[float(x) for x in row] for row in read_csv(IN/(prefix+'_Fuel_kg_h.csv'))[1:]])
    assert a.shape==(504,3) and np.allclose(a[:,:2],b[:,:2]) and np.allclose(a[:,2]*3.6,b[:,2],atol=1e-8)
    checks.append(prefix+': 504 fuel rows, g/s to kg/h readback verified')

# Scientific plots are annotated as a virtual model, not product charts.
gen=np.asarray([[float(x) for x in row] for row in read_csv(IN/'A_Gen80_MG1EST_FullLoad_Q1.csv')[1:]])
eng=np.asarray([[float(x) for x in row] for row in read_csv(IN/'A_Engine15NA80_FullLoad.csv')[1:]])
op=list(csv.DictReader((IN/'A_65kW_OperatingPoints.csv').open(encoding='utf-8-sig')))
fig,ax=plt.subplots(figsize=(8,4.8)); ax.plot(gen[:,1],gen[:,2],label='Virtual generator Q1 limit');ax.plot(gen[:,1],-gen[:,2],label='Virtual generator Q4 limit')
ax.plot(eng[:,0],eng[:,1],color='black',label='Virtual 1.5L engine WOT')
feas=[x for x in op if x['can_reach_65kW']=='1'];ax.plot([float(x['rpm_engine_and_gen_1to1']) for x in feas],[-float(x['engine_torque_Nm']) for x in feas],'ro',label='65kW net Q4 operating points')
ax.set(xlabel='Shaft speed [rpm]',ylabel='Torque [Nm]',title='ASSUMED virtual design - not measured hardware');ax.set_xlim(0,13000);ax.grid(alpha=.2);ax.legend(fontsize=8);fig.tight_layout()
fig.savefig(R/'figures/Virtual_FullLoad.png',dpi=160);plt.close(fig)
e=np.asarray([[float(x) for x in row] for row in read_csv(IN/'A_Gen80_MG1EST_Efficiency_Q4.csv')[1:]])
mask=np.asarray([[float(row[3])] for row in read_csv(IN/'A_Gen80_MG1EST_AllowedMask.csv')[1:] if row[2]=='Q4']).ravel().astype(bool)
fig,ax=plt.subplots(figsize=(8,4.8));sc=ax.scatter(e[mask,2],-e[mask,3],c=e[mask,4],s=45,cmap='viridis',vmin=70,vmax=100)
ax.plot(gen[:,1],gen[:,2],color='black',lw=1,label='Virtual Q4 absolute torque limit')
ax.plot([float(x['rpm_engine_and_gen_1to1']) for x in feas],[float(x['engine_torque_Nm']) for x in feas],'rx',ms=7,label='65kW net')
ax.set(xlabel='Generator speed [rpm]',ylabel='Absolute generator torque [Nm]',title='Virtual Q4 map - only envelope-valid grid points')
fig.colorbar(sc,ax=ax,label='Generator + inverter efficiency [%]');ax.legend(fontsize=8);ax.grid(alpha=.2);fig.tight_layout();fig.savefig(R/'figures/Virtual_Q4_Efficiency.png',dpi=160);plt.close(fig)

wb=Workbook();overview=wb.active;overview.title='START'
for row in [
 ['CRUISE M Range Extender Input Tables','2026-10-06'],
 ['STATUS','A is a virtual draft. REF is a published completed reference, not fully measured.'],
 ['TARGET','1.5L NA, max engine80kW, net bus65kW. Actual matched target maps not found.'],
 ['Image4 Full Load columns','Voltage(V), Speed(rpm), Torque(N*m)'],
 ['Image5 Efficiency columns','Temperature(degC), Voltage(V), Speed(rpm), Torque(N*m), Efficiency(%)'],
 ['GENERATOR SIGN','Q1 positive speed, positive torque. Q4 positive speed, negative torque. Efficiency positive.'],
 ['VOLTAGE','650V source; requires separate conversion. Do not relabel650 as330V.'],
 ['TEMPERATURE','0C is coordinate only; not0C test data. No thermal dependence.'],
 ['ENVELOPE','Efficiency table includes extension points. Use Full Load and AllowedMask to constrain commands.'],
 ['STARTUP','Efficiency grid1000-13000rpm, nonzero torque. PowerLoss alternative contains0rpm and0Nm.'],
 ['Q4 SOURCE','MG1_EST full map estimated; Q4 full load mirror is additional assumption. MG2 Q4 limit in source.'],
 ['INVERTER','Included in EPA EMOT efficiency. Additional DCDC conversion is separate.'],
 ['VALIDATION','Numeric conversion only. No CRUISE M GUI import or simulation run.'],
 ['COPY','Copy numeric columns from sheet matching CSV filename in INDEX. CSV menus depend on installed version.'],
 ['GUIDE','See accompanying Korean Input Guide for settings, formulas, assumptions and sources.'],
 ['A6000rpm','Engine73.258536kW, +/-116.594581Nm, gen+inv92.878238%, net65kW with97% converter and1kWaux.'],
 ]:overview.append(row)
index=wb.create_sheet('INDEX');index.append(['Sheet','CSV filename','Classification','Rows excluding header'])
special={'UI_fields_transcription':'UI_fields','Scalar_presets':'Scalars','A_65kW_OperatingPoints':'A65_operating'}
for p in sorted(IN.glob('*.csv')):
    name=p.stem
    short=name.replace('A_Engine15NA80','A15').replace('REF_Engine20NA','REF20').replace('A_Gen80_MG1EST','AG80').replace('REF_MG1_EST','REF_MG1').replace('FullLoad','FL').replace('Efficiency','ETA').replace('PowerLoss_W_ALTERNATIVE','LossW').replace('AllowedMask','Mask').replace('ClosedThrottle','CT').replace('_matrix','_grid')
    short=special.get(name,short)[:31]
    assert short not in wb.sheetnames
    ws=wb.create_sheet(short);rows=read_csv(p)
    ws.append(rows[0])
    for row in rows[1:]:ws.append([num(v) for v in row])
    classification='VIRTUAL A / ASSUMED' if name.startswith('A_') else 'PUBLIC REF / COMPLETED' if name.startswith('REF_') else 'UI / SETTINGS'
    index.append([short,p.name,classification,len(rows)-1]);index.cell(index.max_row,1).hyperlink=f"#'{short}'!A1"
    ws.sheet_properties.tabColor='F4B183' if name.startswith('A_') else '9DC3E6'
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,(float,int)):cell.number_format='0.000000'
plots=wb.create_sheet('Plots');plots.append(['ASSUMED virtual design charts; not measured product maps'])
for name,anchor in [('Virtual_FullLoad.png','A3'),('Virtual_Q4_Efficiency.png','A43')]:
    im=Image(str(R/'figures'/name));im.width=800;im.height=480;plots.add_image(im,anchor)
for ws in wb:
    ws.freeze_panes='A2'
    if ws.max_row>1 and ws.max_column>1:ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=32
    for j in range(1,ws.max_column+1):
        length=max(len(str(ws.cell(i,j).value or '')) for i in range(1,min(ws.max_row,25)+1))
        ws.column_dimensions[get_column_letter(j)].width=min(max(18,length+2),55)
    if ws.title=='START':ws.column_dimensions['A'].width=28;ws.column_dimensions['B'].width=110
    if ws.title=='Scalars':ws.column_dimensions['E'].width=78
xlsx=OUT/'RangeExtender_CRUISE_M_Inputs_20261006.xlsx';wb.save(xlsx)
reloaded=load_workbook(xlsx,data_only=True)
assert reloaded['AG80_ETA_Q4'].max_row==248
assert reloaded['AG80_ETA_Q4']['E2'].value>1
assert reloaded['A65_operating'].cell(12,11).value==65
checks.append('XLSX reopened; 247Q4 rows, percent units, operating65kW cell verified')

md=(R/'Input_Guide_ko.md').read_text();d=Document();sec=d.sections[0]
sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(1.8)
style=d.styles['Normal'];style.font.name='Malgun Gothic';style.font.size=Pt(9);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Malgun Gothic')
lines=md.splitlines();i=0
while i<len(lines):
    line=lines[i]
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
        rows=[[x.strip() for x in l.strip('|').split('|')] for l in block if not all(set(v.strip())<=set('-: ') for v in l.strip('|').split('|'))]
        tb=d.add_table(rows=0,cols=len(rows[0]));tb.style='Table Grid'
        for row in rows:
            cells=tb.add_row().cells
            for j,v in enumerate(row):cells[j].text=v
        continue
    if line.startswith('# '):d.add_heading(line[2:],0)
    elif line.startswith('## '):d.add_heading(line[3:],1)
    elif line.strip():d.add_paragraph(line)
    i+=1
d.add_heading('가상 후보의 Full load 및 효율 격자',1)
d.add_picture(str(R/'figures/Virtual_FullLoad.png'),width=Cm(16))
d.add_picture(str(R/'figures/Virtual_Q4_Efficiency.png'),width=Cm(16))
docx=OUT/'RangeExtender_CRUISE_M_Input_Guide_20261006.docx';d.save(docx)
assert len(Document(docx).tables)==4
checks.append('DOCX reopened with4 parameter/source/comparison tables and2 annotated plots')
qa={'stage':'R65C_03','checks':checks,'csv_files':len(list(IN.glob('*.csv'))),'xlsx_sheets':len(wb.sheetnames),
    'actual65kW_generator_continuous_rating_verified':False,'CRUISE_M_run':False,'GUI_import_verified':False,
    'existing_inputs_modified':False,'virtual_status':'OPTIONAL_DRAFT_NOT_APPLIED','next':'save workbook, guide and kit; run separate CRUISE M work copy if assumptions adopted'}
(R/'checkpoints/R65C_03.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
manifest=[]
for p in sorted(R.rglob('*')):
    if p.is_file() and p.suffix not in ['.pyc'] and '__pycache__' not in p.parts and p.name!='deliverable_manifest.json':
        manifest.append({'file':str(p.relative_to(R)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(R/'logs/deliverable_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
zpath=OUT/'RangeExtender_CRUISE_M_CSV_Kit_20261006.zip'
with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
    for p in sorted(R.rglob('*')):
        if p.is_file() and '__pycache__' not in p.parts:z.write(p,'RangeExtender_CRUISE_M/'+str(p.relative_to(R)))
    z.write(xlsx,xlsx.name);z.write(docx,docx.name)
with zipfile.ZipFile(zpath) as z:assert z.testzip() is None
print(json.dumps({'checks':qa,'outputs':[{'file':str(p),'bytes':p.stat().st_size} for p in [xlsx,docx,zpath]]},ensure_ascii=False,indent=2))
