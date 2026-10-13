"""Build professor brief and shared editable deck content from checked facts."""
import json, os
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT.parents[1].parent
OUT=Path(os.environ.get('BOLT_REPORT_OUT','/workspace/scratch/a5204a7b57f4/output/bolt_midterm_20261013'))
OUT.mkdir(parents=True,exist_ok=True)
calc=json.loads((ROOT/'calculation_review.json').read_text())
doc=Document()
sec=doc.sections[0]
sec.page_width=Mm(210); sec.page_height=Mm(297)
sec.top_margin=Mm(16); sec.bottom_margin=Mm(16)
sec.left_margin=Mm(18); sec.right_margin=Mm(18)
for name in ['Normal','Title','Heading 1']:
    style=doc.styles[name]
    style.font.name='Noto Sans KR'; style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans KR')
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),'Noto Sans KR')
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),'Noto Sans KR')
    for border in list(style.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    style.paragraph_format.space_after=Pt(4)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(17)
doc.styles['Title'].font.bold=True
for border in list(doc.element.xpath('.//w:pBdr')):border.getparent().remove(border)
doc.styles['Heading 1'].font.size=Pt(11.5)
doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(7)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(3)
doc.core_properties.title='가상 직렬 EREV 연구 방향과 중간보고 범위'
doc.core_properties.author='문진석'
doc.add_paragraph('가상 직렬 EREV 연구 방향과 중간보고 범위',style='Title')
p=doc.add_paragraph('연구 방향 검토   2026년 10월 13일   중간보고 10월 30일')
p.runs[0].font.size=Pt(9.5)
doc.add_paragraph('Bolt BEV의 공개 시험 재현을 기반으로, 유한 주행 미션에서 성능을 유지하는 사양 경계와 발전 기동 제어의 효과를 분석한다. 저SOC 성능 저하가 반드시 발생한다는 전제는 두지 않는다.')
def section(title,text):
    doc.add_paragraph(title,style='Heading 1')
    doc.add_paragraph(text)
section('1 연구 질문','배터리 가용출력·에너지와 발전계 순 DC 출력·응답이 주어졌을 때 요구 주행성능을 유지하는 경계는 어디이며, S0~S4가 경계 근처의 성능 부족과 준비·운전 비용을 얼마나 바꾸는가?')
doc.add_paragraph('2 성능 정의와 평가 조건',style='Heading 1')
table=doc.add_table(rows=1,cols=3)
table.autofit=False
widths=[Mm(24),Mm(64),Mm(86)]
for col,w in zip(table.columns,widths):col.width=w
rows=[['유형','조건','주요 지표'],['출력 제한','상승 부하·80→120 km/h 가속','요구/달성 출력차·부족시간·가속시간'],['에너지 제한','정속 고부하·속도와 고도 조건','목표속도 유지·필요 SOC 여유']]
for n,row in enumerate(rows):
    cells=table.rows[0].cells if n==0 else table.add_row().cells
    for i,val in enumerate(row):
        cells[i].width=widths[i]; cells[i].text=val
        for p in cells[i].paragraphs:
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            for r in p.runs:r.font.size=Pt(10.5);r.bold=n==0
        if n==0:
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EDEFF2');cells[i]._tc.get_or_add_tcPr().append(shd)
section('3 방법과 기대 기여','순간 출력 조건과 유한시간 에너지 조건을 분리해 해석 경계를 도출하고, 차량 모델의 경계 오차와 불확실성 띠를 평가한다. 지속 부하를 주 분석, 가속을 보조 분석으로 둔다. 기동 제어로 회복 가능한 영역과 발전 용량 부족 영역을 구분해 조건부 사양 결정 기준을 제시한다.')
section('4 유지 항목과 수정 사항','BEV 검증, 기존 S0~S4, 주행성능 우선, 26 kW 가정 기준을 유지한다. 40/65 kW와 3P/2P/1P는 가상 사양 민감도다. 휠파워 트리거는 예고 없는 스텝의 최초 부족을 예방할 수 없으므로 인과적 예측과 이상 정보 상한을 구분한다. 실제 Bolt EREV 최적 SOC와 트리거 자체의 독창성은 주장하지 않는다.')
section('5 현재 결과와 중간보고 범위','054 회생 진단은 전압 RMSE 0.663 V·에너지 +3.57%로 해당 시험 조건부 목표를 충족했다. 060 DIS80은 전압 RMSE 1.519 V·에너지 +9.484%로 에너지 ±5% 목표가 미충족이다. 전체 모델의 독립 검증은 진행 중이다. 회생 분배와 방전 손실 경계를 정리한 뒤 입력을 동결하고 066을 평가한다. 10/30에는 검증 진행 결과와 해석 예상 지도, EREV 실행 계획을 발표한다.')
section('6 한계와 지도교수 검토 사항','저SOC BMS·온도별 출력과 발전 과도는 미검증이다. 준정적 수치와 연료 비용은 가정 범위의 예상으로 제한한다. 성능 차이가 작아도 비용 차이는 별도 평가한다. 검토 요청은 경계 중심 문제 정의, 지속 부하 주 분석과 가속 보조 분석, 가상 발전원 사용 범위, 중간보고의 실제 검증 범위다.')
p=doc.add_paragraph('근거  기존 연구 지침 개정3판 13~18쪽 및 DIS80 감사 결과. GM 2017 Volt 설명서 182쪽은 등판 전 충전 여유 확보를 권고한다. 현행 도로구조규칙 제25조의 도로별 표와 예외를 적용한다. 수치·출처·실행계획은 상세 검토서에 수록.')
p.paragraph_format.space_before=Pt(5)
for r in p.runs:r.font.size=Pt(9)
doc.save(OUT/'Professor_Brief_Bolt_EREV_20261013.docx')

repo='https://github.com/hompe-hue/jinseok/blob/main/'
law='https://www.law.go.kr/LSW/lsSideInfoP.do?docCls=jo&joBrNo=00&joNo=0025&lsiSeq=286673&urlMode=lsScJoRltInfoR'
gm='https://www.chevrolet.com/ownercenter/content/dam/gmownercenter/gmna/dynamic/manuals/2017/Chevrolet/Volt/2k17volt1stPrint.pdf'
guide='Bolt_EREV_CRUISE_M_Claude_Guide_20261006.docx 개정3판, 13~18쪽, 파일 버전2'
slides=[]
def slide(title,lead,paragraphs=None,table=None,notes='',foot=''):
    slides.append(dict(title=title,lead=lead,paragraphs=paragraphs or [],table=table,notes=notes,foot=foot))
slide('가상 직렬 EREV의\n주행성능 유지 경계','연구 방향 검토와 10월 30일 중간보고 준비',[
    '문진석  /  2026년 10월 13일',
    'Bolt BEV 공개 시험 재현을 기반으로 한 조건부 사양·제어 비교'],notes='현재 제출 자료는 연구 방향과 중간보고 구성안이다. 미완료 검증과 실행 계획을 실제 결과와 구분한다.')
slide('연구 문제의 수정','성능 저하의 발생 조건을 먼저 정의하고 제어 효과를 평가한다',[
    '저SOC에서 반드시 출력이 무너진다는 전제는 하드웨어와 BMS 근거가 부족하다.',
    '배터리 크기는 가용출력과 하한까지 버티는 시간을 함께 바꾼다.',
    '사양 여유가 충분한 조건도 포함하고, 제어 효과가 입력 불확실성보다 큰지 확인한다.'],notes='Component sizing of a series hybrid electric vehicle through artificial neural network, DOI 10.1016/j.enconman.2022.115300. 출판사 공개 초록/발췌만 확인. 사양과 제어의 결합 자체를 새로운 기여로 주장하지 않는다.\nhttps://www.sciencedirect.com/science/article/pii/S0196890422000966')
slide('연구 질문과 성능 지표','유한 주행 미션의 요구성능 유지 경계와 S0~S4의 효과',[
    '공통 지표  제한 전 요구 휠출력과 달성 출력의 차이, 부족시간·부족에너지',
    '주 분석  정속 고부하의 목표속도 유지와 필요 SOC 여유',
    '보조 분석  상승 부하·80→120 km/h 가속시간 증가율',
    '비용 비교  동등 성능과 종단 저장에너지 조건에서 준비·회복까지 포함'],notes=guide+'\n목표속도 부족 2km/h가 1초 지속하는 구간 기록. 부족>1kW가0.1초 이상인 이벤트. SOC5% 하한. 연구 분류 규칙이며 인증기준 아님. 저SOC 가속은 실차검증이 아닌 민감도 범위.')
slide('BEV 검증의 현재 결과','시험별 조건부 재현 수준과 독립 검증을 구분한다',table=[
    ['비교 항목','현재 결과','판정과 범위'],
    ['054 전압 / 에너지','0.663 V / +3.57%','해당 시험 조건부 목표 충족'],
    ['060 속도 RMSE','0.421 km/h','실측 속도 추종 개선'],
    ['060 전압 / 에너지','1.519 V / +9.484%','에너지 ±5% 목표 미충족'],
    ['독립 시험 066','신규 결과 미확보','입력 동결 후 평가 예정']],
    notes=repo+'research/cruise_m_build_20261011/runs/ANL61905060_DIS80_20261013/RESULT_KO.md\n속도0~3150s, 전압21~3150s, 에너지20~3150s. 원실행마지막0.896995초CSV미수출. 실측10.3122606069kWh, 모델11.2903249669kWh. 전압목표RMSE≤2V,bias≤1V. 독립검증완료를 뜻하지 않음.\n054: '+repo+'research/cruise_m_build_20261011/runs/ANL61905054_regen_20261013/RESULT.md\n20~1144s 조건부 재현, 같은시험초기전압과LV부하사용,모델최대404.50V충전경계미검증. 같은동결variant의전체조건통과아님.',foot='실제 CRUISE M CSV와 ANL 실측 비교  /  모델 전체 검증 진행 중')
slide('순간 출력과 유한시간 에너지 경계','각 원인의 해결 가능성과 필요한 자원이 다르다',[
    '순간 조건  DC 요구가 배터리 가용출력과 발전 순공급의 합 이하',
    '일정 고부하  사용가능 에너지가 순부족 출력 × 지속시간 이상',
    '지속 발전이 부족하면 SOC 여유는 고갈을 늦춘다. 무기한 유지는 불가능하다.',
    '가변 부하는 SOC·V/I·충전수용·손실을 포함한 상태방정식으로 판정'],notes=guide+'\n경계 특례는 발전지연·상승률, 모터토크/rpm한계, 배터리 내부손실을 포함해 확장. 순적산 또는양의부족적분만으로일반가변부하 충분조건을 보장하지 않음.')
rows=[['정속 조건','DC 요구','26 kW','40 kW']]
for d in calc['duration'][:3]:
    label=f"{d['speed_kmh']} km/h · {d['grade_pct']:.0f}%"
    if d['grade_pct']==6:label+='  극한 검토'
    rows.append([label,f"{d['demand_dc_kw']:.1f} kW",f"{d['depletion_minutes']['26']:.1f}분",'순부족 없음' if d['depletion_minutes']['40'] is None else f"{d['depletion_minutes']['40']:.1f}분"])
slide('지속 부하의 조건부 예상','SOC 여유 10%p를 소진하는 시간의 준정적 검산',table=rows,
    notes='calculation_review.json과check_direction_calculations.py. 질량1877.6704kg, 추가질량120kg, 사용가능58.5kWh, 10%p≈5.85kWh, 효율95%, 보조0.4kW가정. 발전은처음부터지정순DC를연속공급. 65kW는세조건에서순부족없으나열디레이팅/초기지연/과도는미반영.\n'+law+'\n1206%는일반고속도로대표조건이아닌별도극한검토조건. 설계속도와운행속도다름. 현행소형차도로120산지5%+예외1%p를고려하므로규정상불가능단정삭제.',
    foot='가정 계산  /  58.5 kWh·추가 120 kg·효율 95%  /  CRUISE M 실행값 아님')
slide('기동 전략의 인과적 한계','현재 휠파워 감지만으로 예고 없는 스텝의 최초 부족을 막을 수 없다',[
    '해석 예  요구 80 kW, 배터리 한계 60 kW, 발전 최대 26 kW',
    '지연 0.5초·상승률 20 kW/s이면 사후 기동에서 1.5초간 5.56 Wh 부족',
    '시작부터 충족하려면 1.5초 앞선 기동 또는 배터리 출력 여유가 필요',
    'S3는 현재·과거 정보만 사용. 미래 이벤트를 아는 Oracle은 상한 비교로 분리'],notes=guide+'\n기존특례와독립적재적분으로5.555555Wh확인. Lmin=td+(Pdem−PB)/R=1.5s. D>Pg이면사전기동만으로해결불가. 선행정보가없으면step전기동은정보누출가능성조사. 미래Driver목표궤적을인과S3입력에주지않음.')
slide('기존 S0부터 S4까지의 정의','기존 이름과 기본 조건을 유지한다',table=[
    ['군','기동 기준','정보와 준비 조건'],
    ['S0','SOC ON 6% / OFF 8%','현재 SOC'],
    ['S1','S0 + 휠 요구 35 kW 초과','현재 제한 전 요구'],
    ['S2','SOC ON 15% / OFF 18%','충전 준비 비용 포함'],
    ['S3','S0 + 인과적 DC 요구 예측','현재·과거, 최대 5초'],
    ['S4','S2와 S3 요청의 OR','여유 확보와 예측 병용']],
    notes=guide+'\nS1임계20/35/50/65별도보정후동결. 공통SOC5%하한. S1파워OFF문턱ON−5kW1s, 최소ON/OFF10/5s. S3필터0.2s,차분0.2s,여유5kW,H=min(5,td+26/R). 40/65변형은Pg로일반화하는별도기록필요. 템플릿PowerMargin칸에35를직접입력하지말고코드·단위확인.',foot='기존 연구 지침 개정3판 13~14쪽  /  비교용 가정이며 최적값 아님')
slide('하드웨어 축과 주장 범위','가상 순 DC 발전원으로 경계 분석을 먼저 완성한다',[
    '기준 26 kW와 40/65 kW 민감도는 순 DC 공급 상한·응답의 가정 축',
    '3P/2P/1P는 용량·저항·허용전류가 함께 변하는 가상 팩. 질량은 별도 패널',
    '25°C를 주 조건으로 고정. 저SOC BMS·저온·열 디레이팅은 불확실성 범위',
    '연료맵은 확인한 정상 운전영역만 사용. 65 kW까지 비례 외삽하지 않음'],notes=repo+'research/bolt_volt_genset_20261006/README_가상EREV.md\n'+repo+'research/range_extender_20261006/RangeExtender_CRUISE_M/Input_Guide_ko.md\nMGA후보제한탐색최대32.48kW,MGB60.56kW는연속정격입증아님. 별도80kW급가상후보순65정속점실현은다른맵/가정이다. Volt파워스플릿과순수seriesEREV구분. 58.5kWh가정은실측용량확정아님.')
slide('대표 실행과 독립 평가','먼저 해석 경계를 확인하고 차량 실행을 단계적으로 늘린다',[
    '대표 미션 3개  여유 충분, 제어에 민감, 검토 제어 모두 미달',
    '차량 주 비교  3개 미션 × S0~S4 = 15개 실행',
    '독립 미션 6개 실행과 경계 근처 지연·상승률 민감도 6개 추가',
    '초기·종단 자원과 준비·회복 비용 통일. 사후 한계 조정으로 효과를 만들지 않음'],notes=guide+'\n27개는중간보고이후첫차량실행의제안규모. 미사용3미션에서S0와사전선정비교군=6,6개민감도별도. 기존1200저차스윕/180느린응답격자삭제하지않고확대후보보존. 대표미션은수치결과보기전에해석여유로선정하고독립미션은튜닝금지. 수명비용은처리량/RMS전류proxy와구분.')
slide('10월 30일까지의 작업','BEV 검증 상태와 해석 예상을 구분해 보고한다',table=[
    ['기간','주요 작업','완료 조건'],
    ['10/14~16','연구 범위·제어·지표 고정','교수님 설명과 가정표'],
    ['10/17~21','회생·방전 손실 경계 정리','식별 근거와 고정 입력'],
    ['10/22~25','조건 충족 시 066 비교','독립 에너지·전압·수지 평가'],
    ['10/26~30','경계 지도·발표 갱신·연습','실행/실측/예상 구분']],
    notes=repo+'research/cruise_m_build_20261011/runs/remaining_review_20261013/RESULT_KO.md\n순회생차이0.036246kWh는총0.978064차이3.71%. 회생정리만으로+9.484→±5목표를보장하지않음. 방전측ABC/EDUneutral겹침숫자미확보시물리보정확정금지. 066미사용조건/평가구간/온도/초기SOC/CSV/log/.proj확인후독립검증. 목표미달시사후완화하지않고보고. 날짜는계획이며실행결과아님.',foot='회생 정리만으로 현재 DC 에너지 오차가 해소된다는 근거는 부족하다')
slide('기대 기여와 논문 구성','해석 경계의 예측력과 조건부 설계 규칙을 제시한다',[
    '공학적 기여  경계 예측오차, 기동으로 회복 가능한 영역, 불확실성 띠',
    '설계 활용  목표 미션에 필요한 발전 순출력과 SOC 여유의 조건부 지도',
    '논문 구성  선행연구, BEV 검증, 해석 조건, 제어 비교, 독립 평가와 한계',
    '지도교수 검토  문제 정의와 주/보조 분석, 가상 발전원 범위, 중간보고 범위'],notes='GM2017Volt설명서인쇄182쪽MountainMode는급경사20분전선택으로reserve준비권고. 실제Volt는동력분기구조로가상직렬플랜트검증근거아님.\n'+gm+'\n연구기여/독창성은선행연구대비및독립결과로확정. 실제BoltEREV최적SOC,트리거최초주장,검증완료과장금지. 성능차이없음≠연료차이없음. 불확실성보다작은효과는불확정.')
(ROOT/'midterm_content.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'brief':str(OUT/'Professor_Brief_Bolt_EREV_20261013.docx'),'slides':len(slides)},ensure_ascii=False))
